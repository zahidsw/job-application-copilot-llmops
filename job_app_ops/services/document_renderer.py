from __future__ import annotations

import html
import re
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from job_app_ops.schemas import GeneratedArtifact

FONT_FAMILY = "Aptos"
CV_BODY_SIZE = 10.5
LETTER_BODY_SIZE = 12
ACCENT_COLOR = "1F4E79"
MUTED_COLOR = "666666"


class DocumentRenderer:
    def render_derivatives(self, artifact: GeneratedArtifact, target: Path) -> list[GeneratedArtifact]:
        if artifact.artifact_type not in {"cv", "motivation_letter"}:
            return []

        docx_path = target.with_suffix(".docx")
        pdf_path = target.with_suffix(".pdf")

        self._render_docx(artifact.content, docx_path, artifact.artifact_type)
        self._render_pdf(artifact.content, pdf_path, artifact.artifact_type)

        return [
            GeneratedArtifact(
                artifact_type=f"{artifact.artifact_type}_docx",
                file_name=docx_path.name,
                content=artifact.content,
                path=str(docx_path),
            ),
            GeneratedArtifact(
                artifact_type=f"{artifact.artifact_type}_pdf",
                file_name=pdf_path.name,
                content=artifact.content,
                path=str(pdf_path),
            ),
        ]

    def _render_docx(self, content: str, target: Path, artifact_type: str) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

        document = Document()
        _configure_docx(document, artifact_type)
        body_size = LETTER_BODY_SIZE if artifact_type == "motivation_letter" else CV_BODY_SIZE
        is_cv = artifact_type == "cv"
        seen_heading1 = False
        paragraph_index = 0

        for block in _parse_blocks(content):
            block_type = block["type"]
            if block_type == "heading1":
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_cv else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(2 if is_cv else 14)
                run = paragraph.add_run(block["text"])
                _format_run(run, size=18 if is_cv else 16, bold=True, color=ACCENT_COLOR if is_cv else "000000")
                seen_heading1 = True
            elif block_type == "heading2":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(10)
                paragraph.paragraph_format.space_after = Pt(4)
                run = paragraph.add_run(block["text"].upper() if is_cv else block["text"])
                _format_run(run, size=10 if is_cv else 13, bold=True, color=ACCENT_COLOR)
                if is_cv:
                    _add_bottom_border(paragraph)
            elif block_type == "heading3":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(1)
                run = paragraph.add_run(block["text"])
                _format_run(run, size=11 if is_cv else 12, bold=True, color="000000")
            elif block_type == "bullet":
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.space_after = Pt(1 if is_cv else 3)
                run = paragraph.add_run(block["text"])
                _format_run(run, size=body_size)
            elif block_type == "numbered":
                paragraph = document.add_paragraph(style="List Number")
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.space_after = Pt(2)
                run = paragraph.add_run(block["text"])
                _format_run(run, size=body_size)
            elif block_type == "paragraph":
                paragraph = document.add_paragraph()
                if is_cv and seen_heading1 and paragraph_index <= 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    color = MUTED_COLOR
                    size = 10.5
                else:
                    color = "000000"
                    size = body_size
                paragraph.paragraph_format.space_after = Pt(3 if is_cv else 8)
                paragraph.paragraph_format.line_spacing = 1.05 if is_cv else 1.15
                run = paragraph.add_run(block["text"])
                _format_run(run, size=size, color=color)
            else:
                if not is_cv:
                    document.add_paragraph("")
            if block_type != "spacer":
                paragraph_index += 1
        document.save(target)

    def _render_pdf(self, content: str, target: Path, artifact_type: str) -> None:
        styles = getSampleStyleSheet()
        body_size = LETTER_BODY_SIZE if artifact_type == "motivation_letter" else CV_BODY_SIZE
        styles.add(ParagraphStyle(name="CvHeading3", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=11, leading=13, spaceAfter=3))
        styles["BodyText"].fontSize = body_size
        styles["BodyText"].leading = body_size + 2
        story: list[object] = []
        bullet_buffer: list[str] = []

        def flush_bullets() -> None:
            nonlocal bullet_buffer
            if not bullet_buffer:
                return
            items = [ListItem(Paragraph(html.escape(item), styles["BodyText"])) for item in bullet_buffer]
            story.append(ListFlowable(items, bulletType="bullet"))
            story.append(Spacer(1, 8))
            bullet_buffer = []

        for block in _parse_blocks(content):
            block_type = block["type"]
            if block_type == "bullet":
                bullet_buffer.append(block["text"])
                continue

            flush_bullets()
            if block_type == "heading1":
                story.append(Paragraph(html.escape(block["text"]), styles["Heading1"]))
            elif block_type == "heading2":
                story.append(Paragraph(html.escape(block["text"]), styles["Heading2"]))
            elif block_type == "heading3":
                story.append(Paragraph(html.escape(block["text"]), styles["CvHeading3"]))
            elif block_type == "paragraph":
                story.append(Paragraph(html.escape(block["text"]), styles["BodyText"]))
            story.append(Spacer(1, 8))

        flush_bullets()
        document = SimpleDocTemplate(str(target), pagesize=A4, leftMargin=0.7 * inch, rightMargin=0.7 * inch, topMargin=0.65 * inch, bottomMargin=0.65 * inch)
        document.build(story)


def _parse_blocks(content: str) -> list[dict[str, str]]:
    blocks: list[dict[str, str]] = []
    first_content_seen = False

    for line in content.splitlines():
        raw_text = re.sub(r"\s+", " ", line).strip()
        text = _strip_inline_markdown(raw_text)
        if not text:
            blocks.append({"type": "spacer", "text": ""})
        elif raw_text.startswith("### "):
            blocks.append({"type": "heading3", "text": _strip_inline_markdown(raw_text[4:].strip())})
        elif text.startswith("# "):
            blocks.append({"type": "heading1", "text": _strip_inline_markdown(raw_text[2:].strip())})
        elif text.startswith("## "):
            blocks.append({"type": "heading2", "text": _strip_inline_markdown(raw_text[3:].strip())})
        elif _is_bold_section_heading(raw_text, text) and first_content_seen:
            blocks.append({"type": "heading2", "text": text.rstrip(":")})
        elif re.match(r"^\*\*[^*]{2,80}\*\*$", raw_text) and not first_content_seen:
            blocks.append({"type": "heading1", "text": text})
        elif re.match(r"^[-*\u2022]\s+", raw_text):
            blocks.append({"type": "bullet", "text": _strip_inline_markdown(re.sub(r"^[-*\u2022]\s+", "", raw_text))})
        elif re.match(r"^\d+[.)]\s+", raw_text):
            blocks.append({"type": "numbered", "text": _strip_inline_markdown(re.sub(r"^\d+[.)]\s+", "", raw_text))})
        else:
            blocks.append({"type": "paragraph", "text": text})
        if text:
            first_content_seen = True
    return blocks


def _is_bold_section_heading(raw_text: str, cleaned_text: str) -> bool:
    if not re.match(r"^\*\*[^*]{2,80}:?\*\*$", raw_text):
        return False
    if cleaned_text.endswith(":"):
        return True
    return cleaned_text.rstrip(":").upper() in {
        "PROFILE",
        "PROFESSIONAL SUMMARY",
        "TARGET FIT",
        "CORE FIT",
        "TECHNICAL SKILLS",
        "PROFESSIONAL EXPERIENCE",
        "EXPERIENCE",
        "EDUCATION",
        "LANGUAGES",
        "CONTACT INFORMATION",
    }


def _configure_docx(document, artifact_type: str) -> None:
    from docx.shared import Inches, Pt

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = FONT_FAMILY
    normal.font.size = Pt(LETTER_BODY_SIZE if artifact_type == "motivation_letter" else CV_BODY_SIZE)
    _set_style_font(normal, FONT_FAMILY)

    for style_name in ("List Bullet", "List Number"):
        if style_name in document.styles:
            style = document.styles[style_name]
            style.font.name = FONT_FAMILY
            style.font.size = normal.font.size
            _set_style_font(style, FONT_FAMILY)


def _set_style_font(style, font_name: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run_properties = style.element.get_or_add_rPr()
    run_fonts = run_properties.rFonts
    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        run_properties.append(run_fonts)
    run_fonts.set(qn("w:ascii"), font_name)
    run_fonts.set(qn("w:hAnsi"), font_name)


def _format_run(run, *, size: float, bold: bool = False, color: str = "000000") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = FONT_FAMILY
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run_properties = run._element.get_or_add_rPr()
    run_fonts = run_properties.rFonts
    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        run_properties.append(run_fonts)
    run_fonts.set(qn("w:ascii"), FONT_FAMILY)
    run_fonts.set(qn("w:hAnsi"), FONT_FAMILY)


def _add_bottom_border(paragraph, color: str = "D9E2EC") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph_properties = paragraph._p.get_or_add_pPr()
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is None:
        border = OxmlElement("w:pBdr")
        paragraph_properties.append(border)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    border.append(bottom)


def _strip_inline_markdown(text: str) -> str:
    cleaned = text.strip()
    cleaned = re.sub(r"\[([^\]]+)\]\((?:mailto:)?([^)]+)\)", lambda match: match.group(1) if match.group(1) == match.group(2) else match.group(1), cleaned)
    cleaned = re.sub(r"`([^`]+)`", r"\1", cleaned)
    cleaned = re.sub(r"\*\*([^*]+)\*\*", r"\1", cleaned)
    cleaned = re.sub(r"__([^_]+)__", r"\1", cleaned)
    cleaned = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip()
